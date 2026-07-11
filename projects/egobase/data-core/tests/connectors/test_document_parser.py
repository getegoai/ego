from pathlib import Path

import pytest

from ego_data_core.connectors.documents import DocumentParser, split_text_into_chunks
from ego_data_core.connectors.filesystem import FilesystemScanner
from ego_data_core.connectors.source_service import SourceService
from ego_data_core.storage.db import get_connection
from ego_data_core.vault import VaultStorage


def test_parse_txt_version_writes_extracted_text_and_chunks(sample_workspace: Path) -> None:
    _scan_workspace(sample_workspace)
    version_id = _version_id_for_title("notes.txt")

    result = DocumentParser().parse_version(version_id)

    assert result.status == "succeeded"
    assert result.evidence_id is not None
    assert result.extracted_text_uri == f"vault://extracted_text/{version_id}.txt"
    assert result.chunk_count == 1
    assert "Ego Data Core sample note" in VaultStorage().read_text(result.extracted_text_uri)

    with get_connection() as connection:
        version = connection.execute(
            "SELECT extracted_text_uri FROM object_versions WHERE version_id = ?",
            (version_id,),
        ).fetchone()
        chunk = connection.execute(
            """
            SELECT *
            FROM evidence_text_chunks
            WHERE evidence_id = ?
            ORDER BY chunk_index
            """,
            (result.evidence_id,),
        ).fetchone()
        fts_count = connection.execute(
            "SELECT COUNT(*) FROM evidence_text_fts WHERE evidence_id = ?",
            (result.evidence_id,),
        ).fetchone()[0]

    assert version["extracted_text_uri"] == result.extracted_text_uri
    assert chunk["chunk_index"] == 0
    assert "safe test data" in chunk["text"]
    assert chunk["char_start"] == 0
    assert chunk["char_end"] > chunk["char_start"]
    assert fts_count == 1


def test_parse_unparsed_versions_continues_after_unsupported_file(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "ok.txt").write_text("This text should still parse.", encoding="utf-8")
    (workspace / "bad.xlsx").write_bytes(b"not really a spreadsheet")
    _scan_workspace(
        workspace,
        scan_policy={"include_globs": ["*.txt", "*.xlsx"]},
    )

    results = DocumentParser().parse_unparsed_versions()

    assert {result.status for result in results} == {"succeeded", "failed"}
    assert _result_for_title(results, "ok.txt").chunk_count == 1
    failed = _result_for_title(results, "bad.xlsx")
    assert failed.status == "failed"
    assert failed.job_id is not None
    assert "Unsupported document type" in failed.error_message

    with get_connection() as connection:
        failed_job = connection.execute(
            "SELECT * FROM jobs WHERE job_id = ?",
            (failed.job_id,),
        ).fetchone()
        error_audit_count = connection.execute(
            "SELECT COUNT(*) FROM audit_logs WHERE event_type = 'error'"
        ).fetchone()[0]
        parsed_versions = connection.execute(
            """
            SELECT COUNT(*)
            FROM object_versions
            WHERE extracted_text_uri IS NOT NULL
            """
        ).fetchone()[0]

    assert failed_job["status"] == "failed"
    assert failed_job["job_type"] == "document_parse"
    assert failed_job["error_message"]
    assert error_audit_count == 1
    assert parsed_versions == 1


def test_chunking_keeps_character_ranges_and_size_bounds() -> None:
    text = "\n\n".join(
        [
            "alpha " * 180,
            "beta " * 180,
            "gamma " * 180,
        ]
    ).strip()

    chunks = split_text_into_chunks(text)

    assert len(chunks) == 3
    assert all(0 < len(chunk.text) <= 1200 for chunk in chunks)
    assert [chunk.char_start for chunk in chunks] == sorted(
        chunk.char_start for chunk in chunks
    )
    assert all(text[chunk.char_start : chunk.char_end] == chunk.text for chunk in chunks)


def test_parse_docx_when_dependency_is_available(tmp_path: Path) -> None:
    pytest.importorskip("docx")
    from docx import Document

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    document = Document()
    document.add_paragraph("DOCX evidence paragraph")
    document.add_paragraph("Second paragraph for parser coverage")
    document.save(workspace / "sample.docx")
    _scan_workspace(
        workspace,
        scan_policy={"include_globs": ["*.docx"]},
    )
    version_id = _version_id_for_title("sample.docx")

    result = DocumentParser().parse_version(version_id)

    assert result.status == "succeeded"
    assert result.chunk_count == 1
    assert "DOCX evidence paragraph" in VaultStorage().read_text(
        result.extracted_text_uri
    )


def test_parse_blank_pdf_when_dependency_is_available(tmp_path: Path) -> None:
    pytest.importorskip("pypdf")
    from pypdf import PdfWriter

    workspace = tmp_path / "workspace"
    workspace.mkdir()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with (workspace / "blank.pdf").open("wb") as file:
        writer.write(file)
    _scan_workspace(
        workspace,
        scan_policy={"include_globs": ["*.pdf"]},
    )
    version_id = _version_id_for_title("blank.pdf")

    result = DocumentParser().parse_version(version_id)

    assert result.status == "succeeded"
    assert result.chunk_count == 0
    assert VaultStorage().read_text(result.extracted_text_uri) == ""


def _scan_workspace(
    workspace: Path,
    *,
    scan_policy: dict[str, object] | None = None,
) -> None:
    source = SourceService().create_source(
        source_type="filesystem",
        name="Parser Test Workspace",
        root_uri=str(workspace),
        scan_policy=scan_policy,
    )
    FilesystemScanner().scan_source(source.source_id)


def _version_id_for_title(title: str) -> str:
    with get_connection() as connection:
        row = connection.execute(
            """
            SELECT ov.version_id
            FROM raw_objects ro
            JOIN object_versions ov ON ov.version_id = ro.current_version_id
            WHERE ro.title = ?
            """,
            (title,),
        ).fetchone()
    assert row is not None
    return row["version_id"]


def _result_for_title(results, title: str):
    version_id = _version_id_for_title(title)
    for result in results:
        if result.version_id == version_id:
            return result
    raise AssertionError(f"Missing parse result for {title}")
