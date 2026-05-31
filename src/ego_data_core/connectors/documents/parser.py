"""Document text extraction and chunking."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import json
from pathlib import Path
import re
from typing import Any

from ego_data_core.audit.service import append_audit_log
from ego_data_core.kernel.id_generator import generate_id
from ego_data_core.storage.db import get_connection, run_migrations, transaction
from ego_data_core.vault.storage import VaultStorage, build_vault_uri


SUPPORTED_DOCUMENT_EXTENSIONS = frozenset({".txt", ".md", ".pdf", ".docx"})
DEFAULT_MIN_CHUNK_CHARS = 800
DEFAULT_MAX_CHUNK_CHARS = 1200


@dataclass(frozen=True)
class DocumentParseResult:
    version_id: str
    evidence_id: str | None
    status: str
    extracted_text_uri: str | None
    chunk_count: int
    error_message: str | None = None
    job_id: str | None = None


@dataclass(frozen=True)
class _VersionContext:
    version_id: str
    raw_object_id: str
    evidence_id: str
    storage_uri: str
    storage_path: Path
    title: str | None
    mime_type: str | None


@dataclass(frozen=True)
class _ChunkDraft:
    text: str
    char_start: int
    char_end: int


class DocumentParserError(RuntimeError):
    """Raised when a document cannot be parsed."""


class UnsupportedDocumentTypeError(DocumentParserError):
    """Raised when the document extension is not supported by v0."""


class MissingDocumentDependencyError(DocumentParserError):
    """Raised when an optional parser dependency is unavailable."""


class DocumentParser:
    def __init__(
        self,
        database_path: str | Path | None = None,
        storage: VaultStorage | None = None,
    ) -> None:
        self.database_path = database_path
        self.storage = storage or VaultStorage()

    def parse_version(self, version_id: str) -> DocumentParseResult:
        run_migrations(self.database_path)
        context = self._load_version_context(version_id)
        try:
            text = _normalize_text(_extract_text(context.storage_path))
            extracted_text_uri = build_vault_uri(
                "extracted_text",
                f"{context.version_id}.txt",
            )
            self.storage.safe_write_text(extracted_text_uri, text)
            chunks = split_text_into_chunks(text)
            self._persist_success(context, extracted_text_uri, chunks)
        except Exception as exc:
            return self._persist_failure(context, exc)

        return DocumentParseResult(
            version_id=context.version_id,
            evidence_id=context.evidence_id,
            status="succeeded",
            extracted_text_uri=extracted_text_uri,
            chunk_count=len(chunks),
        )

    def parse_unparsed_versions(self, *, limit: int | None = None) -> list[DocumentParseResult]:
        run_migrations(self.database_path)
        params: list[Any] = []
        limit_sql = ""
        if limit is not None:
            safe_limit = max(1, int(limit))
            limit_sql = "LIMIT ?"
            params.append(safe_limit)

        with get_connection(self.database_path) as connection:
            rows = connection.execute(
                f"""
                SELECT version_id
                FROM object_versions
                WHERE extracted_text_uri IS NULL
                ORDER BY captured_at ASC, version_id ASC
                {limit_sql}
                """,
                params,
            ).fetchall()

        return [self.parse_version(row["version_id"]) for row in rows]

    def _load_version_context(self, version_id: str) -> _VersionContext:
        with get_connection(self.database_path) as connection:
            row = connection.execute(
                """
                SELECT
                  ov.version_id,
                  ov.raw_object_id,
                  ov.storage_uri,
                  ro.title,
                  ro.mime_type,
                  ev.evidence_id
                FROM object_versions ov
                JOIN raw_objects ro ON ro.raw_object_id = ov.raw_object_id
                LEFT JOIN evidence ev ON ev.version_id = ov.version_id
                WHERE ov.version_id = ?
                ORDER BY ev.created_at ASC, ev.evidence_id ASC
                LIMIT 1
                """,
                (version_id,),
            ).fetchone()

        if row is None:
            raise ValueError(f"Object version not found: {version_id}")
        if row["evidence_id"] is None:
            raise DocumentParserError(f"Object version has no evidence: {version_id}")

        return _VersionContext(
            version_id=row["version_id"],
            raw_object_id=row["raw_object_id"],
            evidence_id=row["evidence_id"],
            storage_uri=row["storage_uri"],
            storage_path=self.storage.resolve_uri(row["storage_uri"]),
            title=row["title"],
            mime_type=row["mime_type"],
        )

    def _persist_success(
        self,
        context: _VersionContext,
        extracted_text_uri: str,
        chunks: list[_ChunkDraft],
    ) -> None:
        now = _utc_now()
        with get_connection(self.database_path) as connection:
            with transaction(connection):
                connection.execute(
                    """
                    UPDATE object_versions
                    SET extracted_text_uri = ?
                    WHERE version_id = ?
                    """,
                    (extracted_text_uri, context.version_id),
                )
                connection.execute(
                    "DELETE FROM evidence_text_fts WHERE evidence_id = ?",
                    (context.evidence_id,),
                )
                connection.execute(
                    "DELETE FROM evidence_text_chunks WHERE evidence_id = ?",
                    (context.evidence_id,),
                )
                for index, chunk in enumerate(chunks):
                    chunk_id = generate_id("chk")
                    connection.execute(
                        """
                        INSERT INTO evidence_text_chunks(
                          chunk_id,
                          evidence_id,
                          chunk_index,
                          text,
                          token_count,
                          language,
                          page_no,
                          char_start,
                          char_end,
                          bbox_json,
                          created_at
                        )
                        VALUES (?, ?, ?, ?, ?, NULL, NULL, ?, ?, NULL, ?)
                        """,
                        (
                            chunk_id,
                            context.evidence_id,
                            index,
                            chunk.text,
                            _estimate_token_count(chunk.text),
                            chunk.char_start,
                            chunk.char_end,
                            now,
                        ),
                    )
                    connection.execute(
                        """
                        INSERT INTO evidence_text_fts(chunk_id, evidence_id, text)
                        VALUES (?, ?, ?)
                        """,
                        (chunk_id, context.evidence_id, chunk.text),
                    )

    def _persist_failure(
        self,
        context: _VersionContext,
        error: Exception,
    ) -> DocumentParseResult:
        message = str(error)
        job_id = self._record_failed_job(context, message)
        append_audit_log(
            "error",
            actor="system",
            target={"type": "object_version", "id": context.version_id},
            payload={
                "component": "document_parser",
                "job_id": job_id,
                "version_id": context.version_id,
                "evidence_id": context.evidence_id,
                "storage_uri": context.storage_uri,
                "error": message,
            },
            database_path=self.database_path,
        )
        return DocumentParseResult(
            version_id=context.version_id,
            evidence_id=context.evidence_id,
            status="failed",
            extracted_text_uri=None,
            chunk_count=0,
            error_message=message,
            job_id=job_id,
        )

    def _record_failed_job(self, context: _VersionContext, message: str) -> str:
        now = _utc_now()
        job_id = generate_id("job")
        with get_connection(self.database_path) as connection:
            connection.execute(
                """
                INSERT INTO jobs(
                  job_id,
                  job_type,
                  status,
                  input_json,
                  output_json,
                  policy_json,
                  error_message,
                  scheduled_at,
                  started_at,
                  finished_at,
                  created_at,
                  updated_at
                )
                VALUES (?, 'document_parse', 'failed', ?, '{}', '{}', ?, ?, ?, ?, ?, ?)
                """,
                (
                    job_id,
                    _to_json(
                        {
                            "version_id": context.version_id,
                            "evidence_id": context.evidence_id,
                            "storage_uri": context.storage_uri,
                        }
                    ),
                    message,
                    now,
                    now,
                    now,
                    now,
                    now,
                ),
            )
            connection.commit()
        return job_id


def split_text_into_chunks(
    text: str,
    *,
    min_chars: int = DEFAULT_MIN_CHUNK_CHARS,
    max_chars: int = DEFAULT_MAX_CHUNK_CHARS,
) -> list[_ChunkDraft]:
    if not text.strip():
        return []

    spans = _paragraph_spans(text)
    chunks: list[_ChunkDraft] = []
    pending_start: int | None = None
    pending_end: int | None = None

    for span_start, span_end in spans:
        if span_end - span_start > max_chars:
            if pending_start is not None and pending_end is not None:
                _append_chunk(chunks, text, pending_start, pending_end)
                pending_start = None
                pending_end = None
            chunks.extend(_split_long_span(text, span_start, span_end, max_chars))
            continue

        if pending_start is None:
            pending_start = span_start
            pending_end = span_end
            continue

        assert pending_end is not None
        candidate = _trimmed_text(text, pending_start, span_end)
        if len(candidate) <= max_chars:
            pending_end = span_end
            continue

        _append_chunk(chunks, text, pending_start, pending_end)
        pending_start = span_start
        pending_end = span_end

    if pending_start is not None and pending_end is not None:
        _append_chunk(chunks, text, pending_start, pending_end)

    if len(chunks) >= 2 and len(chunks[-1].text) < min_chars:
        previous = chunks[-2]
        final = chunks[-1]
        merged_text = _trimmed_text(text, previous.char_start, final.char_end)
        if len(merged_text) <= max_chars:
            chunks[-2:] = [
                _ChunkDraft(
                    text=merged_text,
                    char_start=previous.char_start,
                    char_end=final.char_end,
                )
            ]

    return chunks


def _extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise UnsupportedDocumentTypeError(f"Unsupported document type: {extension or '<none>'}")
    if extension in {".txt", ".md"}:
        return _read_plain_text(path)
    if extension == ".pdf":
        return _extract_pdf_text(path)
    if extension == ".docx":
        return _extract_docx_text(path)
    raise UnsupportedDocumentTypeError(f"Unsupported document type: {extension}")


def _read_plain_text(path: Path) -> str:
    content = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "gb18030", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise MissingDocumentDependencyError("pypdf is required for PDF parsing") from exc

    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text.strip())
    return "\n\n".join(pages)


def _extract_docx_text(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise MissingDocumentDependencyError("python-docx is required for DOCX parsing") from exc

    document = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _normalize_text(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


def _paragraph_spans(text: str) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    for match in re.finditer(r"\S[\s\S]*?(?=\n\s*\n|\Z)", text):
        start, end = _trim_span(text, match.start(), match.end())
        if start < end:
            spans.append((start, end))
    return spans


def _split_long_span(text: str, start: int, end: int, max_chars: int) -> list[_ChunkDraft]:
    chunks: list[_ChunkDraft] = []
    cursor = start
    while cursor < end:
        chunk_end = min(cursor + max_chars, end)
        if chunk_end < end:
            boundary = max(
                text.rfind(" ", cursor, chunk_end),
                text.rfind("\n", cursor, chunk_end),
            )
            if boundary > cursor + max_chars // 2:
                chunk_end = boundary
        _append_chunk(chunks, text, cursor, chunk_end)
        cursor = chunk_end
        while cursor < end and text[cursor].isspace():
            cursor += 1
    return chunks


def _append_chunk(chunks: list[_ChunkDraft], text: str, start: int, end: int) -> None:
    start, end = _trim_span(text, start, end)
    if start < end:
        chunks.append(
            _ChunkDraft(
                text=text[start:end],
                char_start=start,
                char_end=end,
            )
        )


def _trim_span(text: str, start: int, end: int) -> tuple[int, int]:
    while start < end and text[start].isspace():
        start += 1
    while end > start and text[end - 1].isspace():
        end -= 1
    return start, end


def _trimmed_text(text: str, start: int, end: int) -> str:
    start, end = _trim_span(text, start, end)
    return text[start:end]


def _estimate_token_count(text: str) -> int:
    words = text.split()
    if words:
        return len(words)
    return max(1, len(text) // 4)


def _to_json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=True, sort_keys=True, separators=(",", ":"))


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()
